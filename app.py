import os
import sys
import sqlite3
import random
import string
import io
import csv
from datetime import datetime, timedelta
from functools import wraps
from flask import Flask, render_template, request, jsonify, send_file, redirect, url_for, session, flash
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import qrcode

# Ensure UTF-8 output on Windows consoles
try:
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')
except Exception:
    pass

app = Flask(__name__)
app.config['SECRET_KEY'] = 'sorovnoma-portal-super-secret-key-2026'
app.config['PERMANENT_SESSION_LIFETIME'] = timedelta(days=30)
app.config['TEMPLATES_AUTO_RELOAD'] = True

# Admin paroli (buni xohlagan payt o'zgartirish mumkin)
ADMIN_PASSWORD = os.environ.get('ADMIN_PASSWORD', 'admin2026')

DB_PATH = os.path.join(os.path.dirname(__file__), 'database.db')

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db()
    with conn:
        conn.execute('''
            CREATE TABLE IF NOT EXISTS sorovnoma (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp TEXT NOT NULL,
                fingerprint TEXT NOT NULL,
                til TEXT DEFAULT "O'zbek",
                fish TEXT NOT NULL,
                sinf TEXT NOT NULL,
                fanlar TEXT DEFAULT "",
                togaraklar TEXT DEFAULT "",
                iqtidor TEXT DEFAULT "",
                kasb TEXT DEFAULT "",
                startap TEXT DEFAULT "",
                yangi_togaraklar TEXT DEFAULT "",
                takliflar TEXT DEFAULT ""
            )
        ''')
    conn.close()

init_db()

def generate_fingerprint():
    chars = string.ascii_uppercase + string.digits
    suffix = ''.join(random.choices(chars, k=6))
    return f"FP{suffix}"

def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('is_admin'):
            if request.path.startswith('/api/'):
                return jsonify({'status': 'error', 'message': "Avtorizatsiyadan o'tilmagan. Admin huquqi talab qilinadi."}), 401
            return redirect(url_for('admin_login', next=request.path))
        return f(*args, **kwargs)
    return decorated_function

# --- Public Routes (Saytga kirganda faqat so'rovnoma ochiladi) ---

@app.route('/')
def index():
    # Asosiy sahifada FAQAT so'rovnoma ochiladi!
    return render_template('form.html')

@app.route('/sorovnoma')
def survey_redirect():
    return redirect(url_for('index'))

# --- Admin Authentication ---

@app.route('/admin/login', methods=['GET', 'POST'])
def admin_login():
    if session.get('is_admin'):
        return redirect(url_for('admin_table'))

    error = None
    if request.method == 'POST':
        password = request.form.get('password', '').strip()
        if password == ADMIN_PASSWORD:
            session.permanent = True
            session['is_admin'] = True
            next_url = request.args.get('next') or url_for('admin_table')
            return redirect(next_url)
        else:
            error = "Parol noto'g'ri! Qayta urinib ko'ring."

    return render_template('login.html', error=error)

@app.route('/admin/logout')
def admin_logout():
    session.pop('is_admin', None)
    return redirect(url_for('index'))

# --- Protected Admin Pages ---

@app.route('/admin')
@admin_required
def admin_table():
    return render_template('table.html')

@app.route('/admin/stats')
@admin_required
def admin_stats():
    return render_template('stats.html')

@app.route('/stats')
@admin_required
def stats_redirect():
    return redirect(url_for('admin_stats'))

# --- API Routes ---

@app.route('/api/entries', methods=['GET'])
@admin_required
def get_entries():
    query = request.args.get('q', '').strip().lower()
    sinf = request.args.get('sinf', '').strip()
    til = request.args.get('til', '').strip()
    kasb = request.args.get('kasb', '').strip()

    sql = "SELECT * FROM sorovnoma WHERE 1=1"
    params = []

    if query:
        sql += " AND (LOWER(fish) LIKE ? OR LOWER(fingerprint) LIKE ? OR LOWER(takliflar) LIKE ? OR LOWER(iqtidor) LIKE ? OR LOWER(kasb) LIKE ?)"
        params.extend([f"%{query}%", f"%{query}%", f"%{query}%", f"%{query}%", f"%{query}%"])
    if sinf:
        if sinf.endswith('-sinf') or sinf.isdigit() or sinf.endswith('-klass') or sinf.endswith('-класс'):
            grade_num = sinf.split('-')[0]
            sql += " AND (sinf = ? OR sinf LIKE ? OR sinf LIKE ?)"
            params.extend([sinf, f"{grade_num}-%", f"{grade_num}-класс%"])
        else:
            sql += " AND sinf = ?"
            params.append(sinf)
    if til:
        sql += " AND til = ?"
        params.append(til)
    if kasb:
        sql += " AND LOWER(kasb) LIKE ?"
        params.append(f"%{kasb.lower()}%")

    sql += " ORDER BY id DESC"

    conn = get_db()
    rows = conn.execute(sql, params).fetchall()
    conn.close()

    entries = [dict(row) for row in rows]
    return jsonify({'status': 'success', 'count': len(entries), 'data': entries})

# POST: Har kim (o'quvchi) so'rovnoma to'ldirishi uchun ochiq
@app.route('/api/entries', methods=['POST'])
def add_entry():
    data = request.get_json() or {}

    fish = data.get('fish', '').strip()
    sinf = data.get('sinf', '').strip()

    if not fish or not sinf:
        return jsonify({'status': 'error', 'message': "F.I.Sh. va Sinf maydonlarini to'ldirish shart!"}), 400

    # Anti-duplicate / Spam protection: check if same fish and sinf was submitted in last 2 minutes
    conn = get_db()
    recent = conn.execute(
        "SELECT id, timestamp FROM sorovnoma WHERE LOWER(fish) = ? AND sinf = ? ORDER BY id DESC LIMIT 1",
        (fish.lower(), sinf)
    ).fetchone()
    if recent:
        try:
            prev_time = datetime.strptime(recent['timestamp'], '%Y-%m-%d %H:%M:%S')
            diff_secs = (datetime.now() - prev_time).total_seconds()
            if diff_secs < 120:
                conn.close()
                return jsonify({
                    'status': 'error',
                    'message': f"Hurmatli {fish}, sizning anketangiz qabul qilingan! Qayta yuborish talab etilmaydi."
                }), 400
        except Exception:
            pass

    now_str = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    timestamp = data.get('timestamp') or now_str
    fingerprint = data.get('fingerprint') or generate_fingerprint()
    til = data.get('til', "O'zbek").strip()
    fanlar = data.get('fanlar', '')
    if isinstance(fanlar, list):
        fanlar = ", ".join(fanlar)
    togaraklar = data.get('togaraklar', '')
    if isinstance(togaraklar, list):
        togaraklar = ", ".join(togaraklar)
    iqtidor = data.get('iqtidor', '')
    if isinstance(iqtidor, list):
        iqtidor = ", ".join(iqtidor)
    kasb = data.get('kasb', '').strip()
    startap = data.get('startap', '').strip()
    yangi_togaraklar = data.get('yangi_togaraklar', '')
    if isinstance(yangi_togaraklar, list):
        yangi_togaraklar = ", ".join(yangi_togaraklar)
    takliflar = data.get('takliflar', '').strip()

    cursor = conn.cursor()
    cursor.execute('''
        INSERT INTO sorovnoma (timestamp, fingerprint, til, fish, sinf, fanlar, togaraklar, iqtidor, kasb, startap, yangi_togaraklar, takliflar)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (timestamp, fingerprint, til, fish, sinf, fanlar, togaraklar, iqtidor, kasb, startap, yangi_togaraklar, takliflar))
    new_id = cursor.lastrowid
    conn.commit()
    conn.close()

    return jsonify({
        'status': 'success',
        'message': "Ma'lumot muvaffaqiyatli saqlandi!",
        'data': {
            'id': new_id,
            'timestamp': timestamp,
            'fingerprint': fingerprint,
            'til': til,
            'fish': fish,
            'sinf': sinf,
            'fanlar': fanlar,
            'togaraklar': togaraklar,
            'iqtidor': iqtidor,
            'kasb': kasb,
            'startap': startap,
            'yangi_togaraklar': yangi_togaraklar,
            'takliflar': takliflar
        }
    }), 201

@app.route('/api/entries/<int:entry_id>', methods=['GET'])
@admin_required
def get_entry(entry_id):
    conn = get_db()
    row = conn.execute("SELECT * FROM sorovnoma WHERE id = ?", (entry_id,)).fetchone()
    conn.close()
    if not row:
        return jsonify({'status': 'error', 'message': "Yozuv topilmadi"}), 404
    return jsonify({'status': 'success', 'data': dict(row)})

@app.route('/api/entries/<int:entry_id>', methods=['PUT'])
@admin_required
def update_entry(entry_id):
    data = request.get_json() or {}
    conn = get_db()
    row = conn.execute("SELECT * FROM sorovnoma WHERE id = ?", (entry_id,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'status': 'error', 'message': "Yozuv topilmadi"}), 404

    fish = data.get('fish', row['fish']).strip()
    sinf = data.get('sinf', row['sinf']).strip()
    til = data.get('til', row['til']).strip()
    fanlar = data.get('fanlar', row['fanlar'])
    if isinstance(fanlar, list): fanlar = ", ".join(fanlar)
    togaraklar = data.get('togaraklar', row['togaraklar'])
    if isinstance(togaraklar, list): togaraklar = ", ".join(togaraklar)
    iqtidor = data.get('iqtidor', row['iqtidor'])
    if isinstance(iqtidor, list): iqtidor = ", ".join(iqtidor)
    kasb = data.get('kasb', row['kasb']).strip()
    startap = data.get('startap', row['startap']).strip()
    yangi_togaraklar = data.get('yangi_togaraklar', row['yangi_togaraklar'])
    if isinstance(yangi_togaraklar, list): yangi_togaraklar = ", ".join(yangi_togaraklar)
    takliflar = data.get('takliflar', row['takliflar']).strip()

    conn.execute('''
        UPDATE sorovnoma
        SET fish = ?, sinf = ?, til = ?, fanlar = ?, togaraklar = ?, iqtidor = ?, kasb = ?, startap = ?, yangi_togaraklar = ?, takliflar = ?
        WHERE id = ?
    ''', (fish, sinf, til, fanlar, togaraklar, iqtidor, kasb, startap, yangi_togaraklar, takliflar, entry_id))
    conn.commit()
    conn.close()

    return jsonify({'status': 'success', 'message': "Yozuv yangilandi!"})

@app.route('/api/entries/<int:entry_id>', methods=['DELETE'])
@admin_required
def delete_entry(entry_id):
    conn = get_db()
    conn.execute("DELETE FROM sorovnoma WHERE id = ?", (entry_id,))
    conn.commit()
    conn.close()
    return jsonify({'status': 'success', 'message': "Yozuv o'chirildi!"})

@app.route('/api/clear', methods=['POST'])
@admin_required
def clear_all():
    conn = get_db()
    conn.execute("DELETE FROM sorovnoma")
    conn.commit()
    conn.close()
    return jsonify({'status': 'success', 'message': "Barcha ma'lumotlar tozalandi!"})

# --- Statistics API ---

@app.route('/api/stats', methods=['GET'])
@admin_required
def get_stats():
    conn = get_db()
    rows = conn.execute("SELECT * FROM sorovnoma").fetchall()
    conn.close()

    total = len(rows)
    sinflar = {}
    tillari = {}
    kasblar = {}
    yangi_togaraklar = {}
    fanlar = {}
    iqtidorlar = {}

    for r in rows:
        s = r['sinf'] or 'Noma\'lum'
        sinflar[s] = sinflar.get(s, 0) + 1

        t = r['til'] or 'O\'zbek'
        tillari[t] = tillari.get(t, 0) + 1

        k = (r['kasb'] or '').strip()
        if k:
            kasblar[k] = kasblar.get(k, 0) + 1

        yt = r['yangi_togaraklar'] or ''
        for item in [x.strip() for x in yt.split(',') if x.strip()]:
            yangi_togaraklar[item] = yangi_togaraklar.get(item, 0) + 1

        fn = r['fanlar'] or ''
        for item in [x.strip() for x in fn.split(',') if x.strip()]:
            fanlar[item] = fanlar.get(item, 0) + 1

        iq = r['iqtidor'] or ''
        for item in [x.strip() for x in iq.split(',') if x.strip()]:
            iqtidorlar[item] = iqtidorlar.get(item, 0) + 1

    top_kasblar = sorted(kasblar.items(), key=lambda x: x[1], reverse=True)[:10]
    top_yangi_togaraklar = sorted(yangi_togaraklar.items(), key=lambda x: x[1], reverse=True)[:10]
    top_fanlar = sorted(fanlar.items(), key=lambda x: x[1], reverse=True)[:10]
    top_iqtidorlar = sorted(iqtidorlar.items(), key=lambda x: x[1], reverse=True)[:8]

    return jsonify({
        'total': total,
        'sinflar': sinflar,
        'tillari': tillari,
        'top_kasblar': top_kasblar,
        'top_yangi_togaraklar': top_yangi_togaraklar,
        'top_fanlar': top_fanlar,
        'top_iqtidorlar': top_iqtidorlar
    })

# --- Export Routes ---

@app.route('/export/excel')
@admin_required
def export_excel():
    sinf = request.args.get('sinf', '').strip()
    conn = get_db()
    if sinf:
        if sinf.endswith('-sinf') or sinf.isdigit() or sinf.endswith('-klass') or sinf.endswith('-класс'):
            grade_num = sinf.split('-')[0]
            rows = conn.execute("SELECT timestamp, fingerprint, til, fish, sinf, fanlar, togaraklar, iqtidor, kasb, startap, yangi_togaraklar, takliflar FROM sorovnoma WHERE (sinf = ? OR sinf LIKE ? OR sinf LIKE ?) ORDER BY id ASC", (sinf, f"{grade_num}-%", f"{grade_num}-класс%")).fetchall()
        else:
            rows = conn.execute("SELECT timestamp, fingerprint, til, fish, sinf, fanlar, togaraklar, iqtidor, kasb, startap, yangi_togaraklar, takliflar FROM sorovnoma WHERE sinf = ? ORDER BY id ASC", (sinf,)).fetchall()
        sheet_title = f"{sinf} sinf"
        filename = f"Sorovnoma_{sinf}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    else:
        rows = conn.execute("SELECT timestamp, fingerprint, til, fish, sinf, fanlar, togaraklar, iqtidor, kasb, startap, yangi_togaraklar, takliflar FROM sorovnoma ORDER BY id ASC").fetchall()
        sheet_title = "Barcha natijalar"
        filename = f"Sorovnoma_35maktab_Barchasi_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    headers = [
        "Sana va vaqt", "Fingerprint", "Til", "F.I.Sh.", "Sinf",
        "Qiziqadigan fanlar", "Togaraklar", "Iqtidor sohasi", "Kelajak kasbi",
        "Startap loyihasi", "Yangi togaraklar taklifi", "Qoshimcha takliflar"
    ]

    header_fill = PatternFill(start_color="107C41", end_color="107C41", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    alt_fill = PatternFill(start_color="F9FBF9", end_color="F9FBF9", fill_type="solid")

    for r_idx, row in enumerate(rows, start=2):
        row_data = list(row)
        ws.append(row_data)
        is_alt = (r_idx % 2 == 0)
        for col_num in range(1, len(row_data) + 1):
            cell = ws.cell(row=r_idx, column=col_num)
            cell.font = Font(name="Arial", size=10)
            cell.border = thin_border
            if col_num in [1, 2, 3, 5]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align
            if is_alt:
                cell.fill = alt_fill
        ws.row_dimensions[r_idx].height = 22

    col_widths = {
        'A': 20, 'B': 14, 'C': 12, 'D': 32, 'E': 14,
        'F': 35, 'G': 35, 'H': 35, 'I': 26, 'J': 30,
        'K': 35, 'L': 40
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )

@app.route('/export/csv')
@admin_required
def export_csv():
    sinf = request.args.get('sinf', '').strip()
    conn = get_db()
    if sinf:
        if sinf.endswith('-sinf') or sinf.isdigit() or sinf.endswith('-klass') or sinf.endswith('-класс'):
            grade_num = sinf.split('-')[0]
            rows = conn.execute("SELECT timestamp, fingerprint, til, fish, sinf, fanlar, togaraklar, iqtidor, kasb, startap, yangi_togaraklar, takliflar FROM sorovnoma WHERE (sinf = ? OR sinf LIKE ? OR sinf LIKE ?) ORDER BY id ASC", (sinf, f"{grade_num}-%", f"{grade_num}-класс%")).fetchall()
        else:
            rows = conn.execute("SELECT timestamp, fingerprint, til, fish, sinf, fanlar, togaraklar, iqtidor, kasb, startap, yangi_togaraklar, takliflar FROM sorovnoma WHERE sinf = ? ORDER BY id ASC", (sinf,)).fetchall()
        filename = f"Sorovnoma_{sinf}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    else:
        rows = conn.execute("SELECT timestamp, fingerprint, til, fish, sinf, fanlar, togaraklar, iqtidor, kasb, startap, yangi_togaraklar, takliflar FROM sorovnoma ORDER BY id ASC").fetchall()
        filename = f"Sorovnoma_35maktab_Barchasi_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    conn.close()

    headers = [
        "Sana va vaqt", "Fingerprint", "Til", "F.I.Sh.", "Sinf",
        "Qiziqadigan fanlar", "Togaraklar", "Iqtidor sohasi", "Kelajak kasbi",
        "Startap loyihasi", "Yangi togaraklar taklifi", "Qoshimcha takliflar"
    ]

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for r in rows:
        writer.writerow(list(r))

    mem = io.BytesIO()
    mem.write(output.getvalue().encode('utf-8-sig'))
    mem.seek(0)

    return send_file(
        mem,
        mimetype="text/csv",
        as_attachment=True,
        download_name=filename
    )

@app.route('/api/qr')
def get_qr():
    host = request.host
    url = f"https://{host}/" if not host.startswith("localhost") and not host.startswith("127.0.0.1") else f"http://{host}/"
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_L,
        box_size=10,
        border=2,
    )
    qr.add_data(url)
    qr.make(fit=True)
    img = qr.make_image(fill_color="#0F9D58", back_color="white")
    
    img_io = io.BytesIO()
    img.save(img_io, 'PNG')
    img_io.seek(0)
    return send_file(img_io, mimetype='image/png')

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    print(f"============================================================")
    print(f"   SO'ROVNOMA VA NATIJALAR PORTALI ISHGA TUSHDI!            ")
    print(f"   Asosiy So'rovnoma: http://localhost:{port}/              ")
    print(f"   Yopiq Admin Paneli: http://localhost:{port}/admin        ")
    print(f"   Admin Paroli: {ADMIN_PASSWORD}                           ")
    print(f"============================================================")
    app.run(host='0.0.0.0', port=port, debug=False)
